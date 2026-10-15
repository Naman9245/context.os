"""Build small real PDF / Word / PowerPoint files for tests."""

from __future__ import annotations

from pathlib import Path


def make_pdf(path: Path, text: str = "", *, title: str | None = None) -> Path:
    import pymupdf

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = pymupdf.open()
    page = doc.new_page()
    y = 72
    for line in text.splitlines() or [""]:
        page.insert_text((72, y), line, fontsize=11)
        y += 16
    if title:
        doc.set_metadata({"title": title})
    doc.save(path)
    doc.close()
    return path


def make_scanned_pdf(path: Path, text: str) -> Path:
    """A PDF that is only a picture of text - no text layer, like a phone scan."""
    import pymupdf

    path.parent.mkdir(parents=True, exist_ok=True)
    src = pymupdf.open()
    page = src.new_page()
    y = 72
    for line in text.splitlines():
        page.insert_text((72, y), line, fontsize=14)
        y += 24
    image = page.get_pixmap(dpi=150).tobytes("png")
    scan = pymupdf.open()
    scan.new_page(width=page.rect.width, height=page.rect.height).insert_image(
        page.rect, stream=image)
    scan.save(path)
    return path


def make_docx(path: Path, text: str) -> Path:
    import docx

    path.parent.mkdir(parents=True, exist_ok=True)
    d = docx.Document()
    for line in text.splitlines():
        d.add_paragraph(line)
    d.save(path)
    return path


def make_pptx(path: Path, title: str, body: str = "") -> Path:
    import pptx

    path.parent.mkdir(parents=True, exist_ok=True)
    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = body
    deck.save(path)
    return path
