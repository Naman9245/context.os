"""Create a folder of realistic, messy student files for demos and manual testing.

    .venv/bin/python scripts/make_demo_files.py ~/contextos-demo/Downloads

Deadlines are written relative to today, so the demo always has upcoming ones.
"""

from __future__ import annotations

import shutil
import sys
from datetime import date, timedelta
from pathlib import Path


def pdf(path: Path, text: str) -> None:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    y = 72
    for line in text.splitlines():
        page.insert_text((72, y), line, fontsize=11)
        y += 16
    doc.save(path)
    doc.close()


def scanned_pdf(path: Path, text: str) -> None:
    """Only a picture of the text, no text layer - what a phone scanner app produces."""
    import pymupdf

    src = pymupdf.open()
    page = src.new_page()
    y = 72
    for line in text.splitlines():
        page.insert_text((72, y), line, fontsize=14)
        y += 24
    image = page.get_pixmap(dpi=150).tobytes("png")
    scan = pymupdf.open()
    scan.new_page(width=page.rect.width, height=page.rect.height).insert_image(
        page.rect, stream=image)
    scan.save(path)


def docx(path: Path, text: str) -> None:
    import docx as python_docx

    d = python_docx.Document()
    for line in text.splitlines():
        d.add_paragraph(line)
    d.save(path)


def pptx(path: Path, title: str, body: str) -> None:
    import pptx as python_pptx

    deck = python_pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = body
    deck.save(path)


def main(out: Path) -> None:
    out.mkdir(parents=True, exist_ok=True)
    today = date.today()

    def due(days: int) -> str:
        return (today + timedelta(days=days)).strftime("%d/%m/%Y")

    def due_words(days: int) -> str:
        d = today + timedelta(days=days)
        return f"{d.day} {d.strftime('%B %Y')}"

    pdf(out / "DBMS_Assignment_2.pdf",
        "Database Management Systems (BCS403)\nAssignment 2: Transactions and Concurrency\n"
        "1. Explain ACID properties with examples.\n2. Draw the precedence graph for schedule S.\n"
        f"Last date of submission: {due(9)}")
    pdf(out / "Module3_final(2).pdf",
        "BCS403 - Database Management Systems\nModule 3: Normalization\n"
        "Functional dependencies, 1NF, 2NF, 3NF and BCNF.\nLossless join decomposition.")
    shutil.copy(out / "Module3_final(2).pdf", out / "Module3_final(2) (1).pdf")  # a duplicate
    # An assignment with no deadline written in it: Khushi asks when it's due.
    pdf(out / "OS_Assignment_3.pdf",
        "BCS303 - Operating Systems\nAssignment 3: Memory Management\n"
        "Q1. Explain paging with a neat diagram.\nQ2. Compare FIFO and LRU page replacement.\n"
        "Q3. What is thrashing? How can it be prevented?")
    pdf(out / "os question paper 2024.pdf",
        "BCS303 - Operating Systems\nThird Semester B.E. Examination, Dec 2024\n"
        "Time: 3 hrs   Max. Marks: 100\nAnswer any FIVE full questions.\n"
        "1. Explain the process state diagram.")
    pdf(out / "IMG_20261012_0001.pdf",
        "RV College of Engineering\nDepartment of Computer Science\n"
        "Course: Design and Analysis of Algorithms (BCS401)\nAssignment - II\n"
        "Q1. Prove that merge sort runs in O(n log n).\nQ2. Solve 0/1 knapsack using DP.\n"
        f"Submit hand-written answers by {due_words(20)}.")
    pdf(out / "doc_17.pdf",
        "Dayananda Sagar University\nInternal Assessment Test 1\nSubject: Computer Networks\n"
        "Semester: V   Max Marks: 50   Time: 1.5 hrs\nAnswer any two questions from each part.\n"
        "1. Explain the TCP three-way handshake.")
    pdf(out / "Notice_exam_fee.pdf",
        "NOTICE\nAll students are hereby informed that the semester examination fee\n"
        f"must be paid on or before {due(5)}.\nLate fee applies after the due date.\nRegistrar")
    pptx(out / "Week4_normalisation.pptx", "DBMS - Normalisation",
         "1NF, 2NF, 3NF\nDecomposition\nBCNF")
    docx(out / "CN Lab Manual.docx",
         "Computer Networks Laboratory\nLab Manual\nExperiment 1: Simulate a 3-node network "
         "in NS2\nExperiment 2: Implement CRC")
    # A phone scan: needs OCR (sudo apt install tesseract-ocr) to be read at all.
    scanned_pdf(out / "CamScanner 10-14-2026 10.30.pdf",
                "Operating Systems (BCS303)\nLab Record - Experiment 4\n"
                f"CPU Scheduling: Round Robin\nSubmit on or before {due_words(22)}")
    print(f"wrote {len(list(out.iterdir()))} files to {out}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(Path(sys.argv[1]).expanduser())
